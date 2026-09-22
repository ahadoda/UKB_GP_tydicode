from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
import argparse
import json
from pathlib import Path
import re

import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
SOURCE_VERSION = "UK Biobank lookups and mappings version 4 (June 2023)"


@dataclass(frozen=True)
class TraitRule:
    group: str
    trait: str
    expected_units: str
    include: str
    exclude: str = ""
    core: str = ""
    allow_qualitative_core: bool = False


QUALITATIVE = re.compile(
    r"\b(normal|abnormal|raised|high|low|borderline|target|centile|category|"
    r"positive|negative|elevated|increased|decreased|impaired|declined|refused|"
    r"not indicated|screening|monitoring range)\b|\[d\]|[<>]"
)


TRAITS = [
    TraitRule("anthropometry", "body_weight", "kg; verify source unit", r"\bo/e\s*-\s*(?:body )?weight\b|\bbody weight(?: measurement| measured| level)?\b|\bweight (?:measurement|measured|in kg)\b", r"birth|fetal|foetal|baby|infant|child|symptom|gain|loss|diet|ideal|target|predicted|dry weight|molecular", r"^(?:o/e\s*-\s*)?(?:body )?weight(?: measurement)?$|body weight measured|weight in kg"),
    TraitRule("anthropometry", "body_height", "cm or m; verify source unit", r"\bo/e\s*-\s*(?:body )?height\b|\bbody height(?: measurement| measured| level)?\b|\bheight (?:measurement|measured|in cm|in metres?)\b|\bstature measurement\b", r"fundal|heel|shoe|loss of height|short stature|predicted|expected|parental", r"^(?:o/e\s*-\s*)?(?:body )?height(?: measurement)?$|body height measured|height in (?:cm|metres?)"),
    TraitRule("anthropometry", "body_mass_index", "kg/m2", r"\bbody mass index\b|\bbmi\b|quetelet", r"target|advice|education|review|assessment declined", r"^(?:bmi\s*-\s*)?body mass index(?:\s*-\s*observation)?$|^body mass index$"),
    TraitRule("anthropometry", "waist_circumference", "cm; verify source unit", r"\bwaist circumference\b|\babdominal circumference\b", r"declined|target|fetal|foetal", r"^(?:baseline )?(?:waist|abdominal) circumference$"),
    TraitRule("anthropometry", "hip_circumference", "cm; verify source unit", r"\bhip circumference\b", r"declined|target", r"^(?:baseline )?hip circumference$"),
    TraitRule("anthropometry", "waist_hip_ratio", "ratio", r"\bwaist(?: to|[- ])hip ratio\b|\bwaist:hip ratio\b", r"target|declined", r"waist(?: to|[- ])hip ratio|waist:hip ratio", True),
    TraitRule("haemodynamics", "systolic_blood_pressure", "mmHg", r"systolic (?:blood pressure|bp)|(?:blood pressure|bp).*systolic", r"ventricular|dysfunction|target|goal|difference|drop|response|control|education", r"reading|measurement|level|average|mean|home|ambulatory|standing|sitting|lying|o/e"),
    TraitRule("haemodynamics", "diastolic_blood_pressure", "mmHg", r"diastolic (?:blood pressure|bp)|(?:blood pressure|bp).*diastolic", r"ventricular|dysfunction|target|goal|difference|drop|response|control|education", r"reading|measurement|level|average|mean|home|ambulatory|standing|sitting|lying|o/e"),
    TraitRule("haemodynamics", "blood_pressure", "mmHg", r"\bblood pressure (?:reading|measurement|level)\b|\bo/e\s*-\s*blood pressure\b", r"systolic|diastolic|target|goal|control|education|declined", r"^(?:o/e\s*-\s*)?blood pressure(?: reading| measurement)?$"),
    TraitRule("glycaemia", "blood_glucose", "mmol/L; verify source unit", r"\b(?:blood|serum|plasma) (?:fasting |random |post[- ]?prandial |\d+[- ]?(?:minute|hour) )?glucose (?:level|result|measurement)\b|\boral glucose tolerance (?:test|result)\b|\bogtt\b", r"urine|csf|cerebrospinal|monitor|strip|device|target|education|self measurement|ketone", r"level|result|measurement|tolerance test"),
    TraitRule("glycaemia", "hba1c", "mmol/mol or %; determine from code and value", r"hba1c|glycat(?:ed|ed) h[ae]moglobin|glycosylated h[ae]moglobin", r"target|goal|education|review|poor control|good control|monitoring range|diagnostic reference range", r"level|fraction|dcct aligned|ifcc aligned|ifcc standardised"),
    TraitRule("glycaemia", "insulin", "mU/L or pmol/L; verify source unit", r"\b(?:serum|plasma|fasting) insulin (?:level|measurement|assay)\b|\binsulin assay\b", r"drug|injection|therapy|dose|antibody|receptor|diabetes|resistance", r"level|measurement|assay"),
    TraitRule("glycaemia", "c_peptide", "nmol/L or pmol/L; verify source unit", r"\bc[- ]?peptide (?:level|measurement|assay)\b|\bserum c[- ]?peptide\b|\bplasma c[- ]?peptide\b", r"antibody|therapy|education", r"level|measurement|assay"),
    TraitRule("lipids", "total_cholesterol", "mmol/L", r"\b(?:serum|plasma|blood|total) cholesterol(?: level| measurement)?\b", r"hdl|ldl|vldl|ratio|family history|target|diet|drug|screening", r"^(?:serum|plasma|blood|total) cholesterol(?: level| measurement)?$"),
    TraitRule("lipids", "hdl_cholesterol", "mmol/L", r"\bhdl cholesterol(?: level| measurement)?\b|\bhigh[- ]density lipoprotein cholesterol(?: level| measurement)?\b", r"non[- ]?hdl|ratio|electrophores|deficiency|disease|target", r"level|measurement"),
    TraitRule("lipids", "ldl_cholesterol", "mmol/L", r"\bldl cholesterol(?: level| measurement)?\b|\blow[- ]density lipoprotein cholesterol(?: level| measurement)?\b", r"ratio|electrophores|deficiency|disease|target|hyperchol", r"level|measurement"),
    TraitRule("lipids", "triglycerides", "mmol/L", r"\b(?:serum|plasma|blood) triglycerides?(?: level| measurement)?\b|\btriglyceride level\b", r"ratio|family history|fh:|target|diet|drug", r"^(?:serum|plasma|blood) triglycerides?(?: level| measurement)?$"),
    TraitRule("lipids", "apolipoprotein_a", "g/L; verify analyte and source unit", r"\b(?:serum|plasma) apolipoprotein a(?:[- ]?i|1) level\b", r"deficiency|variant|genotyp|antibody", r"level"),
    TraitRule("lipids", "apolipoprotein_b", "g/L; verify source unit", r"\b(?:serum|plasma) apolipoprotein b(?:-100)? level\b", r"deficiency|variant|genotyp|familial|defective", r"level"),
    TraitRule("lipids", "lipoprotein_a", "nmol/L or mg/dL; do not convert without an assay-specific factor", r"\b(?:serum|plasma) lipoprotein\s*\(a\)(?: level| measurement)?\b", r"antibody|phenotype|genotype", r"level|measurement"),
    TraitRule("renal", "creatinine", "umol/L; verify source unit", r"\b(?:serum|plasma) creatinine(?: level| measurement)?\b", r"phosphokinase|kinase|ratio|clearance|target|urine", r"^(?:serum|plasma) creatinine(?: level| measurement)?$"),
    TraitRule("renal", "egfr", "mL/min/1.73m2; verify equation", r"\bestimated glomerular filtration rate\b|\begfr\b|\bglomerular filtration rate (?:level|calculated|estimated|calculation)\b", r"ckd |category|testing|with gfr|tubular|maximal|target|radioisotope|clearance", r"estimated|calculated|calculation|level"),
    TraitRule("renal", "albumin_creatinine_ratio", "mg/mmol; verify source unit", r"\burine (?:micro)?albumin[: /-]?creatinine ratio\b|\balbumin[: /-]?creatinine ratio\b", r"testing not indicated|target|diabetes mellitus with", r"ratio", True),
    TraitRule("renal", "urea", "mmol/L; verify source unit", r"\b(?:serum|plasma|blood) urea (?:level|measurement)\b", r"urine|ratio|cycle|breath|clearance|nitrogen", r"level|measurement"),
    TraitRule("renal", "cystatin_c", "mg/L; verify source unit", r"\b(?:serum|plasma|blood) cystatin c (?:level|measurement)\b", r"egfr|glomerular filtration", r"level|measurement"),
    TraitRule("liver", "alanine_aminotransferase", "U/L", r"\b(?:serum|plasma) alanine (?:amino)?transferase level\b|\balanine transaminase level\b|\balt(?:/sgpt)? (?:serum )?level\b", r"ratio|target|disorder", r"level"),
    TraitRule("liver", "aspartate_aminotransferase", "U/L", r"\b(?:serum|plasma) aspartate (?:amino)?transferase level\b|\baspartate transaminase level\b|\bast (?:serum )?level\b", r"ratio|target|disorder", r"level"),
    TraitRule("liver", "gamma_glutamyl_transferase", "U/L", r"\bgamma[- ]glutamyl (?:transferase|transpeptidase)(?: level)?\b|\bggt (?:level|gamma)\b", r"cycle|disorder|target", r"level|transferase$|transpeptidase$"),
    TraitRule("liver", "alkaline_phosphatase", "U/L", r"\b(?:serum|plasma) alkaline phosphatase(?: level| measurement)?\b", r"placental|bone[- ]specific|isoenzyme|leucocyte|neutrophil|prostatic|intestinal|heat stable", r"^(?:serum|plasma) alkaline phosphatase(?: level| measurement)?$"),
    TraitRule("liver", "albumin", "g/L; verify source unit", r"\b(?:serum|plasma) albumin(?: level| measurement)?\b", r"urine|ratio|creatinine|globulin|dialysis|ascites", r"^(?:serum|plasma) albumin(?: level| measurement)?$"),
    TraitRule("liver", "bilirubin", "umol/L; distinguish total and direct", r"\b(?:serum|plasma) (?:total |direct |conjugated )?bilirubin(?: level| measurement)?\b", r"urine|transcutaneous|neonatal|foetal|fetal", r"level|measurement"),
    TraitRule("inflammation", "c_reactive_protein", "mg/L", r"\b(?:serum|plasma) c[- ]?reactive protein(?: level| measurement)?\b|\bc[- ]?reactive protein level\b|\bcrp level\b", r"target|antibody|screening", r"level|measurement"),
    TraitRule("metabolic", "urate", "umol/L; verify source unit", r"\b(?:serum|plasma) (?:urate|uric acid)(?: level| measurement)?\b", r"urine|stone|gout|target|diet", r"^(?:serum|plasma) (?:urate|uric acid)(?: level| measurement)?$"),
    TraitRule("metabolic", "calcium", "mmol/L; distinguish corrected and uncorrected", r"\b(?:serum|plasma) (?:adjusted |corrected )?calcium(?: level| measurement)?\b", r"urine|ionised|ionized|score|coronary|diet|supplement", r"level|measurement|^(?:serum|plasma) (?:adjusted |corrected )?calcium$"),
    TraitRule("metabolic", "phosphate", "mmol/L; verify source unit", r"\b(?:serum|plasma) (?:inorganic )?phosphate(?: level| measurement)?\b", r"urine|alkaline|buffer|supplement", r"level|measurement"),
    TraitRule("metabolic", "vitamin_d", "nmol/L; prefer total 25-hydroxyvitamin D", r"\b(?:serum|plasma) (?:total )?(?:25[- ]?hydroxy)?vitamin d(?:2|3)?(?: level| measurement)?\b", r"allergy|supplement|therapy|deficiency|rickets|1,25|dihydroxy", r"25[- ]?hydroxy"),
    TraitRule("thyroid", "tsh", "mIU/L; verify source unit", r"\b(?:serum |plasma )?(?:tsh|thyroid[- ]stimulating hormone)(?: level| measurement)?\b", r"releasing hormone|suppression therapy|deficiency|resistance|secreting|overproduction|receptor|antibody|binding site|blood spot|\d+ minute", r"^(?:serum |plasma )?tsh(?: level| measurement)?$|^tsh\s*-\s*thyroid stim(?:ulating)?\.? hormone(?: \(& level\))?$|^thyroid[- ]stimulating hormone level$"),
    TraitRule("thyroid", "free_t4", "pmol/L; verify source unit", r"\b(?:serum|plasma )?(?:free t4|free thyroxine)(?: level| measurement)?\b", r"index|uptake|ratio|therapy", r"level|measurement"),
    TraitRule("thyroid", "free_t3", "pmol/L; verify source unit", r"\b(?:serum|plasma )?(?:free t3|free triiodothyronine)(?: level| measurement)?\b", r"uptake|ratio|therapy", r"level|measurement"),
    TraitRule("thyroid", "thyroid_peroxidase_antibody", "IU/mL or kIU/L; verify assay", r"\bthyroid peroxidase antibod(?:y|ies)(?: level| concentration| measurement)?\b", r"negative history|family history", r"level|concentration|measurement"),
    TraitRule("thyroid", "thyroglobulin_antibody", "IU/mL or kIU/L; verify assay", r"\b(?:anti[- ]?)?thyroglobulin antibod(?:y|ies)(?: level| concentration| measurement)?\b", r"negative history|family history", r"level|concentration|measurement"),
    TraitRule("haematology", "white_blood_cell_count", "10^9/L", r"\b(?:total )?white (?:blood )?cell count(?: level| measurement)?\b|\bwbc count\b", r"differential|urine|csf|cerebrospinal|declined", r"count|level|measurement"),
    TraitRule("haematology", "mean_corpuscular_volume", "fL", r"\bmean (?:cell|corpuscular) volume(?: level| measurement)?\b|\bmcv(?: level| measurement)?\b", r"platelet|reticulocyte", r"volume|level|measurement"),
    TraitRule("haematology", "red_cell_distribution_width", "%", r"\bred (?:blood )?cell distribut(?:ion)? width(?: level| measurement)?\b|\brdw(?: level| measurement)?\b", r"target", r"width|level|measurement"),
    TraitRule("haematology", "haemoglobin", "g/L; verify source unit", r"\b(?:haemoglobin|hemoglobin|hb) estimation(?: \(& level\))?$|^(?:hb\s*-\s*)?haemoglobin (?:concentration|level)$", r"not estimated|requested|sample sent|glycat|glycosyl|fetal|foetal|carboxy|methaemoglobin|oxyhaemoglobin|urine|plasma|variant|electrophoresis|mean cell|mchc|unstable|alteration", r"estimation|concentration|level"),
    TraitRule("haematology", "platelet_count", "10^9/L", r"\bplatelet count(?: level| measurement| observation)?\b|\bplt\s*-\s*platelet count\b", r"reticulated|low|decreased|thrombocyt", r"count|level|measurement|observation"),
]


# These parent concepts have analyte-specific synonyms in the lookup, but their
# full meaning spans several tests. They are unsafe for a specific numeric trait.
TRAIT_CODE_EXCLUSIONS = {
    "tsh": {"442..", ".442."},
}


MEDICATION_RULES = {
    "glucose_lowering": lambda row: row[6] == "Drugs Used In Diabetes" and row[5] in {"Antidiabetic Drugs", "Insulin"},
    "lipid_lowering": lambda row: row[6] == "Lipid-Regulating Drugs",
    "weight_management": lambda row: row[6] == "Obesity",
    "antihypertensive": lambda row: row[4] in {
        "Angiotensin-Converting Enzyme Inhibitors",
        "Angiotensin-II Receptor Antagonists",
        "Beta-Adrenoceptor Blocking Drugs",
        "Calcium-Channel Blockers",
        "Centrally-Acting Antihypertensive Drugs",
        "Vasodilator Antihypertensive Drugs",
        "Thiazides And Related Diuretics",
    },
}


PUBLIC_NAMES = {
    "body_weight": "Body weight",
    "body_height": "Body height",
    "body_mass_index": "Body mass index (BMI)",
    "waist_circumference": "Waist circumference",
    "hip_circumference": "Hip circumference",
    "blood_pressure": "Blood pressure (generic)",
    "systolic_blood_pressure": "Systolic blood pressure",
    "diastolic_blood_pressure": "Diastolic blood pressure",
    "blood_glucose": "Blood glucose",
    "hba1c": "HbA1c",
    "insulin": "Insulin",
    "c_peptide": "C-peptide",
    "total_cholesterol": "Total cholesterol",
    "hdl_cholesterol": "HDL cholesterol",
    "ldl_cholesterol": "LDL cholesterol",
    "triglycerides": "Triglycerides",
    "apolipoprotein_a": "Apolipoprotein A-I",
    "apolipoprotein_b": "Apolipoprotein B",
    "lipoprotein_a": "Lipoprotein(a)",
    "creatinine": "Creatinine",
    "egfr": "Estimated glomerular filtration rate (eGFR)",
    "albumin_creatinine_ratio": "Urine albumin-creatinine ratio",
    "urea": "Urea",
    "cystatin_c": "Cystatin C",
    "alanine_aminotransferase": "Alanine aminotransferase (ALT)",
    "aspartate_aminotransferase": "Aspartate aminotransferase (AST)",
    "gamma_glutamyl_transferase": "Gamma-glutamyl transferase (GGT)",
    "alkaline_phosphatase": "Alkaline phosphatase (ALP)",
    "albumin": "Serum albumin",
    "bilirubin": "Bilirubin",
    "c_reactive_protein": "C-reactive protein (CRP)",
    "urate": "Urate",
    "calcium": "Calcium",
    "phosphate": "Phosphate",
    "vitamin_d": "25-hydroxyvitamin D",
    "tsh": "Thyroid-stimulating hormone (TSH)",
    "free_t4": "Free T4",
    "free_t3": "Free T3",
    "thyroid_peroxidase_antibody": "Thyroid peroxidase antibody",
    "thyroglobulin_antibody": "Thyroglobulin antibody",
    "white_blood_cell_count": "White blood cell count",
    "mean_corpuscular_volume": "Mean corpuscular volume (MCV)",
    "red_cell_distribution_width": "Red cell distribution width (RDW)",
    "haemoglobin": "Haemoglobin",
    "platelet_count": "Platelet count",
    "glucose_lowering": "Glucose-lowering medicines",
    "lipid_lowering": "Lipid-lowering medicines",
    "antihypertensive": "Antihypertensive medicines",
    "weight_management": "Weight-management medicines",
}


PUBLIC_USE = {
    "anthropometry": "Body size and adiposity",
    "haemodynamics": "Blood-pressure measurements",
    "glycaemia": "Glycaemia and diabetes monitoring",
    "lipids": "Lipid and cardiovascular-risk biomarkers",
    "renal": "Kidney function and kidney damage",
    "liver": "Liver biomarkers",
    "inflammation": "Systemic inflammation",
    "metabolic": "Other cardiometabolic biomarkers",
    "thyroid": "Thyroid function and autoimmunity",
    "haematology": "Blood count and biological-age inputs",
    "medication": "Cardiometabolic prescriptions",
}


def public_recommended_use(row: pd.Series) -> str:
    """Give a short, code-level instruction for the reader-facing list."""
    trait = row["trait"]
    coding_system = row["coding_system"]
    description = str(row["description"])
    text = description.casefold()
    name = PUBLIC_NAMES.get(trait, trait.replace("_", " ").title())

    if row["trait_group"] == "medication":
        if coding_system == "bnf":
            return f"DEFAULT - prefix-match to identify {name.lower()} prescriptions."
        return f"DEFAULT - exact-match to identify {name.lower()} prescriptions."

    if trait == "blood_pressure":
        return (
            "SUPPORTING - use only for combined or non-component blood-pressure "
            "records; do not treat as systolic or diastolic without field context."
        )

    if trait in {"systolic_blood_pressure", "diastolic_blood_pressure"}:
        contexts = [
            ("standing", "standing"),
            ("sitting", "sitting"),
            ("lying", "lying"),
            ("24 hour", "24-hour ambulatory average"),
            ("day interval", "daytime ambulatory average"),
            ("night interval", "night-time ambulatory average"),
            ("home", "home measurement"),
            ("ambulatory", "ambulatory measurement"),
        ]
        for token, label in contexts:
            if token in text:
                return f"CONTEXT - use only when the analysis includes {label} blood pressure."
        if "average" in text or "mean" in text:
            return "CONTEXT - use for an averaged blood-pressure value, not a single reading."
        return f"DEFAULT - use as a direct numeric {name.lower()} measurement."

    if trait == "blood_glucose":
        if "fasting" in text:
            return "CONTEXT - use for fasting glucose analyses."
        if "random" in text:
            return "CONTEXT - use for random glucose analyses."
        if "post-prandial" in text or "after evening meal" in text:
            return "CONTEXT - use for post-prandial glucose analyses."
        if "during night" in text:
            return "CONTEXT - use only for overnight glucose measurements."
        if "baseline" in text:
            return "CONTEXT - use as the baseline value in a timed glucose test."
        timed = re.search(r"(\d+)\s*(minute|hour)", text)
        if timed:
            return (
                f"CONTEXT - use as the {timed.group(1)}-{timed.group(2)} value in a "
                "timed or post-load glucose analysis."
            )
        if "tolerance" in text or "ogtt" in text:
            return "CONTEXT - use for oral glucose-tolerance testing; retain test timing."
        return "DEFAULT - use for glucose analyses when fasting or test timing is not required."

    if trait == "hba1c":
        if "ifcc" in text:
            return "UNIT-SPECIFIC - use as IFCC HbA1c (normally mmol/mol); verify the recorded unit."
        if "dcct" in text:
            return "UNIT-SPECIFIC - use as DCCT HbA1c (normally %); verify the recorded unit."
        return "DEFAULT - use for HbA1c, but confirm the unit before combining records."

    if trait == "egfr":
        if "cystatin c" in text:
            return "CONTEXT - use for cystatin-C-based eGFR; do not mix with creatinine eGFR without a plan."
        if "creatinine" in text:
            return "CONTEXT - use for creatinine-based eGFR; retain the equation where available."
        if "african american" in text:
            return "CONTEXT - legacy race-adjusted MDRD eGFR; analyse separately or exclude per protocol."
        if "modification of diet" in text:
            return "CONTEXT - use for MDRD eGFR; retain the equation in harmonisation."
        if "epidemiology collaboration" in text or "ckd-epi" in text:
            return "CONTEXT - use for CKD-EPI eGFR; retain the equation in harmonisation."
        return "DEFAULT - use for eGFR when the equation is not required; verify units."

    if trait == "bilirubin":
        if "conjugated" in text or "direct" in text:
            return "CONTEXT - use for direct/conjugated bilirubin, not total bilirubin."
        if "total" in text:
            return "DEFAULT - use for total bilirubin analyses."
        return "DEFAULT - use for bilirubin only when the source field confirms the fraction measured."

    if trait == "calcium":
        if "adjusted" in text or "corrected" in text:
            return "CONTEXT - use for albumin-adjusted/corrected calcium."
        return "DEFAULT - use for unadjusted total calcium; do not mix with corrected calcium without a plan."

    if trait == "vitamin_d":
        if re.search(r"vitamin d\s*[23]", text):
            return "CONTEXT - use for the named D2 or D3 fraction, not total 25-hydroxyvitamin D."
        return "DEFAULT - use for total 25-hydroxyvitamin D; verify assay and units."

    if "baseline" in text:
        return f"CONTEXT - use only when a baseline {name.lower()} measurement is intended."

    return f"DEFAULT - use as a direct numeric {name.lower()} measurement; verify the recorded unit."


def compile_rules() -> list[tuple[TraitRule, re.Pattern[str], re.Pattern[str] | None, re.Pattern[str]]]:
    compiled = []
    for rule in TRAITS:
        compiled.append(
            (
                rule,
                re.compile(rule.include),
                re.compile(rule.exclude) if rule.exclude else None,
                re.compile(rule.core or rule.include),
            )
        )
    return compiled


def select_traits(workbook: openpyxl.Workbook) -> list[pd.DataFrame]:
    frames: list[pd.DataFrame] = []
    compiled = compile_rules()
    sheet_specs = {
        "read_v2_lkp": {"coding_system": "read_v2", "code": 0, "term": 2, "meta": 1},
        "read_ctv3_lkp": {"coding_system": "ctv3", "code": 0, "term": 1, "meta": 2},
    }

    for sheet_name, spec in sheet_specs.items():
        matches: dict[tuple[str, str], dict[str, object]] = {}
        worksheet = workbook[sheet_name]
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if row[spec["code"]] is None or row[spec["term"]] is None:
                continue
            code = str(row[spec["code"]]).strip()
            term = str(row[spec["term"]]).strip()
            lowered = term.casefold()
            for rule, include, exclude, core in compiled:
                if code in TRAIT_CODE_EXCLUSIONS.get(rule.trait, set()):
                    continue
                if not include.search(lowered) or (exclude and exclude.search(lowered)):
                    continue
                key = (rule.trait, code)
                is_core = bool(core.search(lowered)) and (
                    rule.allow_qualitative_core or not QUALITATIVE.search(lowered)
                )
                record = matches.setdefault(
                    key,
                    {
                        "trait_id": f"{rule.group}.{rule.trait}.{spec['coding_system']}",
                        "trait_group": rule.group,
                        "trait": rule.trait,
                        "coding_system": spec["coding_system"],
                        "source_dataset": "gp_clinical",
                        "source_field": "read_2" if spec["coding_system"] == "read_v2" else "read_3",
                        "match_type": "exact",
                        "case_sensitive": "true",
                        "code": code,
                        "description": term,
                        "matched_terms": [],
                        "selection_tier": "extended_review",
                        "value_type": "numeric_candidate",
                        "expected_units": rule.expected_units,
                        "source_sheet": sheet_name,
                        "source_metadata": "",
                        "review_status": "candidate_requires_clinical_review",
                    },
                )
                record["matched_terms"].append(term)
                if is_core:
                    record["selection_tier"] = "core"
                    record["description"] = term
                metadata = row[spec["meta"]]
                if metadata is not None and not record["source_metadata"]:
                    record["source_metadata"] = str(metadata).strip()

        rows = []
        for record in matches.values():
            record["matched_terms"] = " | ".join(dict.fromkeys(record["matched_terms"]))[:1000]
            rows.append(record)
        sheet_frame = pd.DataFrame(rows).sort_values(
            ["trait_group", "trait", "selection_tier", "code"]
        )
        for _trait_id, trait_frame in sheet_frame.groupby("trait_id", sort=False):
            frames.append(trait_frame.copy())
    return frames


def select_medications(workbook: openpyxl.Workbook) -> tuple[list[pd.DataFrame], dict[str, set[str]]]:
    worksheet = workbook["bnf_lkp"]
    group_rows: dict[str, dict[str, dict[str, str]]] = {
        group: {} for group in MEDICATION_RULES
    }
    group_legacy_bnf: dict[str, set[str]] = {group: set() for group in MEDICATION_RULES}

    def legacy_bnf_category(presentation_code: str) -> str:
        if len(presentation_code) < 7:
            return ""
        subparagraph = presentation_code[6]
        if subparagraph.isdigit():
            subparagraph = subparagraph.zfill(2)
        return (
            f"{presentation_code[0:2]}.{presentation_code[2:4]}."
            f"{presentation_code[4:6]}.{subparagraph}"
        )

    for source in worksheet.iter_rows(min_row=2, values_only=True):
        row = ["" if value is None else str(value).strip() for value in source]
        if not row[0]:
            continue
        for group, selector in MEDICATION_RULES.items():
            if not selector(row):
                continue
            prefix = row[0][:9]
            legacy_code = legacy_bnf_category(row[0])
            if legacy_code:
                group_legacy_bnf[group].add(legacy_code)
            group_rows[group].setdefault(
                prefix,
                {
                    "trait_id": f"medication.{group}.bnf",
                    "trait_group": "medication",
                    "trait": group,
                    "coding_system": "bnf",
                    "source_dataset": "gp_scripts",
                    "source_field": "bnf_code",
                    "match_type": "prefix",
                    "case_sensitive": "true",
                    "code": prefix,
                    "description": row[3] or row[2] or row[1],
                    "matched_terms": row[3] or row[2] or row[1],
                    "selection_tier": "core",
                    "value_type": "prescription_exposure",
                    "expected_units": "not applicable",
                    "source_sheet": "bnf_lkp",
                    "source_metadata": " | ".join(value for value in row[4:8] if value),
                    "review_status": "candidate_requires_pharmacology_review",
                },
            )

    frames = [pd.DataFrame(group_rows[group].values()).sort_values("code") for group in MEDICATION_RULES]
    return frames, group_legacy_bnf


def map_read_drugs(
    workbook: openpyxl.Workbook, group_legacy_bnf: dict[str, set[str]]
) -> list[pd.DataFrame]:
    descriptions: dict[str, str] = {}
    for row in workbook["read_v2_drugs_lkp"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None and row[1] is not None:
            descriptions.setdefault(str(row[0]).strip(), str(row[1]).strip())

    rows_by_group: dict[str, dict[str, dict[str, str]]] = {
        group: {} for group in group_legacy_bnf
    }
    for row in workbook["read_v2_drugs_bnf"].iter_rows(min_row=2, values_only=True):
        if row[0] is None or row[1] is None:
            continue
        read_code = str(row[0]).strip()
        bnf_code = str(row[1]).strip()
        for group, categories in group_legacy_bnf.items():
            if bnf_code not in categories:
                continue
            rows_by_group[group].setdefault(
                read_code,
                {
                    "trait_id": f"medication.{group}.read_v2_drug",
                    "trait_group": "medication",
                    "trait": group,
                    "coding_system": "read_v2_drug",
                    "source_dataset": "gp_scripts",
                    "source_field": "read_2",
                    "match_type": "exact",
                    "case_sensitive": "true",
                    "code": read_code,
                    "description": descriptions.get(read_code, ""),
                    "matched_terms": descriptions.get(read_code, ""),
                    "selection_tier": "core",
                    "value_type": "prescription_exposure",
                    "expected_units": "not applicable",
                    "source_sheet": "read_v2_drugs_bnf",
                    "source_metadata": f"Mapped BNF code: {bnf_code}",
                    "review_status": "candidate_requires_pharmacology_review",
                },
            )
    frames = []
    for group in group_legacy_bnf:
        if rows_by_group[group]:
            frames.append(pd.DataFrame(rows_by_group[group].values()).sort_values("code"))
    return frames


def output_path(frame: pd.DataFrame) -> Path:
    first = frame.iloc[0]
    if first["trait_group"] == "medication":
        folder = ROOT / "codelists" / "medications" / first["trait"]
    else:
        folder = ROOT / "codelists" / "clinical_traits" / first["trait_group"] / first["trait"]
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f"{first['coding_system']}.csv"


def write_public_codelists(master: pd.DataFrame) -> None:
    """Write the simple, reader-facing lists used from the GitHub front page."""
    public_root = ROOT / "phenotypes"
    index_rows = []
    for (group, trait, coding_system), data in master.groupby(
        ["trait_group", "trait", "coding_system"], sort=True
    ):
        core = data[data["selection_tier"].eq("core")]
        selected = core if not core.empty else data
        if group == "medication":
            folder = public_root / "medications"
            prefix = "gp_bnf" if coding_system == "bnf" else "gp_read2drugs"
        else:
            folder = public_root / group
            prefix = "gp_read2" if coding_system == "read_v2" else "gp_read3"
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{prefix}_{trait}.txt"
        simple = selected[["code", "description"]].drop_duplicates().sort_values("code")
        simple["recommended_use"] = selected.loc[simple.index].apply(
            public_recommended_use, axis=1
        )
        simple.to_csv(path, sep="\t", index=False, lineterminator="\n")
        index_rows.append(
            {
                "research_use": PUBLIC_USE[group],
                "phenotype": PUBLIC_NAMES.get(trait, trait.replace("_", " ").title()),
                "ukb_table": selected.iloc[0]["source_dataset"],
                "ukb_field": selected.iloc[0]["source_field"],
                "coding_system": coding_system,
                "matching": selected.iloc[0]["match_type"],
                "expected_units": selected.iloc[0]["expected_units"],
                "code_count": len(simple),
                "codelist": path.relative_to(ROOT).as_posix(),
            }
        )
    pd.DataFrame(index_rows).sort_values(
        ["research_use", "phenotype", "coding_system"]
    ).to_csv(public_root / "index.csv", index=False, lineterminator="\n")


def build(source: Path) -> None:
    if not source.exists():
        raise FileNotFoundError(source)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    frames = select_traits(workbook)
    medication_frames, prefixes = select_medications(workbook)
    frames.extend(medication_frames)
    frames.extend(map_read_drugs(workbook, prefixes))
    workbook.close()

    frames = [frame for frame in frames if not frame.empty]
    manifest = []
    for frame in frames:
        path = output_path(frame)
        frame.to_csv(path, index=False, lineterminator="\n")
        first = frame.iloc[0]
        manifest.append(
            {
                "trait_id": first["trait_id"],
                "trait_group": first["trait_group"],
                "trait": first["trait"],
                "coding_system": first["coding_system"],
                "source_dataset": first["source_dataset"],
                "source_field": first["source_field"],
                "match_type": first["match_type"],
                "case_sensitive": first["case_sensitive"],
                "expected_units": first["expected_units"],
                "core_rows": int(frame["selection_tier"].eq("core").sum()),
                "extended_review_rows": int(frame["selection_tier"].eq("extended_review").sum()),
                "total_rows": len(frame),
                "path": path.relative_to(ROOT).as_posix(),
                "review_status": first["review_status"],
            }
        )

    catalogue = ROOT / "catalogue"
    catalogue.mkdir(parents=True, exist_ok=True)
    master = pd.concat(frames, ignore_index=True).sort_values(
        ["trait_group", "trait", "coding_system", "selection_tier", "code"]
    )
    master.to_csv(catalogue / "clinical_trait_codelists.csv", index=False, lineterminator="\n")
    write_public_codelists(master)
    pd.DataFrame(manifest).sort_values(
        ["trait_group", "trait", "coding_system"]
    ).to_csv(catalogue / "trait_manifest.csv", index=False, lineterminator="\n")
    rule_rows = [
        {
            "trait_group": rule.group,
            "trait": rule.trait,
            "include_regex": rule.include,
            "exclude_regex": rule.exclude,
            "code_exclusions": " | ".join(sorted(TRAIT_CODE_EXCLUSIONS.get(rule.trait, set()))),
            "core_regex": rule.core,
            "expected_units": rule.expected_units,
            "selection_method": "terminology_description",
        }
        for rule in TRAITS
    ]
    rule_rows.extend(
        {
            "trait_group": "medication",
            "trait": trait,
            "include_regex": "",
            "exclude_regex": "",
            "code_exclusions": "",
            "core_regex": "",
            "expected_units": "not applicable",
            "selection_method": "named_bnf_hierarchy_categories",
        }
        for trait in MEDICATION_RULES
    )
    pd.DataFrame(rule_rows).to_csv(
        catalogue / "selection_rules.csv", index=False, lineterminator="\n"
    )
    summary = (
        master.groupby(["trait_group", "trait", "coding_system", "selection_tier"])
        .size()
        .rename("row_count")
        .reset_index()
    )
    summary.to_csv(catalogue / "review_summary.csv", index=False, lineterminator="\n")
    metadata = {
        "source_file": source.name,
        "source_version": SOURCE_VERSION,
        "method": "Rule-based candidate selection from local lookup descriptions and BNF hierarchy",
        "trait_count": int(master[["trait_group", "trait"]].drop_duplicates().shape[0]),
        "codelist_count": len(frames),
        "row_count": len(master),
        "core_row_count": int(master["selection_tier"].eq("core").sum()),
        "review_status": "All outputs require domain review before analysis",
    }
    (catalogue / "build_metadata.json").write_text(
        json.dumps(metadata, indent=2) + "\n", encoding="utf-8"
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source",
        type=Path,
        default=ROOT.parent / "primarycare_codings" / "all_lkps_maps_v4.xlsx",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build(args.source.resolve())
